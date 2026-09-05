import os

import pytest
from docx import Document
from pydantic import ValidationError
from pypdf import PdfReader

from docflow.config import Settings
from docflow.documents import InvalidTemplate, convert_pdf, render_docx
from docflow.schemas import ProposalContent


def test_rounding_uses_decimal_per_line(content):
    content["line_items"] = [
        {"description": "Half day", "quantity": "0.50", "unit_cents": 101},
        {"description": "Half day", "quantity": "0.50", "unit_cents": 101},
    ]
    proposal = ProposalContent.model_validate(content)
    assert proposal.total_cents == 102


@pytest.mark.parametrize(
    "change",
    [
        {"title": ""},
        {"summary": "Short"},
        {"currency": "XYZ"},
        {"unexpected": True},
        {"client_name": "invalid\x00character"},
        {"line_items": []},
        {"line_items": [{"description": "Work", "quantity": "-1", "unit_cents": 100}]},
        {"line_items": [{"description": "Work", "quantity": "1.001", "unit_cents": 100}]},
        {"line_items": [{"description": "Work", "quantity": "1", "unit_cents": 100.5}]},
    ],
)
def test_invalid_inputs_are_rejected(content, change):
    content.update(change)
    with pytest.raises(ValidationError):
        ProposalContent.model_validate(content)


def test_repeating_rows_optional_section_and_xml_escaping(content, tmp_path):
    content["client_name"] = "Research & Development <team>"
    content["assumptions"] = ""
    output = tmp_path / "proposal.docx"
    render_docx(
        Settings().template_path, output, ProposalContent.model_validate(content), 3, "test-ref"
    )
    doc = Document(output)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    assert content["client_name"] in text
    assert "Assumptions" not in text
    for item in content["line_items"]:
        assert item["description"] in text
    assert "EUR 15,900.00" in text
    assert not any(marker in text for marker in ["{{", "{%", "Investment Committee", "Lorem ipsum"])
    assert len(doc.tables[-1].rows) == len(content["line_items"]) + 2


def test_undefined_template_field_fails_loudly(content, tmp_path):
    template = tmp_path / "bad.docx"
    doc = Document()
    doc.add_paragraph("{{ missing_client_field }}")
    doc.save(template)
    with pytest.raises(InvalidTemplate):
        render_docx(
            template, tmp_path / "output.docx", ProposalContent.model_validate(content), 1, "test"
        )


def test_malformed_archive_is_rejected(content, tmp_path):
    template = tmp_path / "bad.docx"
    template.write_bytes(b"not-a-zip")
    with pytest.raises(InvalidTemplate):
        render_docx(
            template, tmp_path / "output.docx", ProposalContent.model_validate(content), 1, "test"
        )


@pytest.mark.conversion
@pytest.mark.skipif(
    not os.getenv("RUN_CONVERSION_TESTS"),
    reason="Set RUN_CONVERSION_TESTS=1 with LibreOffice installed",
)
def test_real_pdf_contains_scope_and_total(content, tmp_path):
    settings = Settings()
    output = tmp_path / "proposal.docx"
    render_docx(
        settings.template_path, output, ProposalContent.model_validate(content), 1, "test-ref"
    )
    pdf = convert_pdf(output, tmp_path, settings.libreoffice_bin)
    reader = PdfReader(pdf)
    text = " ".join(p.extract_text() for p in reader.pages)
    assert 1 <= len(reader.pages) <= 3
    assert "Meridian Supply" in text
    assert "15,900.00" in text
    assert "Frontend implementation" in text
